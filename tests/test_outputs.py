"""Built outputs: page counts, fonts, structure, and agreement between calculations, workbook, PDF and slides."""
import re
import shutil
import subprocess

import openpyxl
import pymupdf
import pytest
from pptx import Presentation

from mpr import theme as T
from mpr.config import SNAPSHOT_DIR

from conftest import norm

EDITIONS = ["2026-08", "2026-09"]


def pdf_text(path):
    return norm(" ".join(p.get_text() for p in pymupdf.open(path)))


def deck_text(path):
    prs = Presentation(path)
    return [norm(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)) for s in prs.slides]


def expected_strings(r):
    """Every headline figure that must agree across outputs, formatted by the single shared formatter."""
    out = {"port": T.pct(r.port_return), "bench": T.pct(r.bench_return), "active": T.pp(r.active_return)}
    for t, row in r.holdings.iterrows():
        out[f"ret_{t}"], out[f"contrib_{t}"] = T.pct(row["total_ret"]), T.pp(row["contrib"])
        out[f"endw_{t}"] = T.wt(row["end_w"])
    return {k: norm(v) for k, v in out.items()}


@pytest.mark.parametrize("eid", EDITIONS)
def test_pdf_is_exactly_two_pages_with_embedded_fonts(built, eid):
    doc = pymupdf.open(built[eid]["pdf"])
    assert len(doc) == 2
    fonts = {f[3] for p in doc for f in p.get_fonts()}
    assert fonts and all("InterTabular" in f for f in fonts), fonts
    assert all(f[1] == "ttf" for p in doc for f in p.get_fonts())
    sizes = {round(s["size"], 1) for p in doc for b in p.get_text("dict")["blocks"] if b["type"] == 0 for l in b["lines"] for s in l["spans"]}
    assert min(sizes) >= 9.0, sizes  # footnotes 9 pt, body 10.5 pt, nothing smaller
    assert 10.5 in sizes


@pytest.mark.parametrize("eid", EDITIONS)
def test_pdf_page_content_and_labels(built, results, eid):
    doc = pymupdf.open(built[eid]["pdf"])
    p1, p2 = norm(doc[0].get_text()), norm(doc[1].get_text())
    r = results[eid]
    for needle in (r.label, "Independent educational research sample".upper(), "Hypothetical", "Market overview", "Facts.", "Interpretation.", f"{r.cutoff:%B %-d, %Y} close", "CAD"):
        assert needle in p1, needle
    for needle in ("Portfolio and benchmark detail", "Equity implications", "Fixed-income implications", "Maintain or rebalance", "Counterargument", "Downside risks", "What to monitor", "Method and sources", "Page 2 of 2"):
        assert needle in p2, needle
    assert len(doc[0].get_images()) == 1 and len(doc[1].get_images()) == 1  # performance chart; contribution chart
    assert "CIBC" not in pdf_text(built[eid]["pdf"])


@pytest.mark.parametrize("eid", EDITIONS)
def test_pdf_links_and_source_references(built, results, eid):
    doc = pymupdf.open(built[eid]["pdf"])
    links = [l["uri"] for p in doc for l in p.get_links() if "uri" in l]
    assert any("bankofcanada.ca" in u for u in links) and any("federalreserve.gov" in u for u in links)
    txt = pdf_text(built[eid]["pdf"])
    for n in range(1, 7):
        assert f"[{n}]" in txt


@pytest.mark.parametrize("eid", EDITIONS)
def test_figures_agree_between_calculations_pdf_and_slides(built, results, eid):
    r = results[eid]
    exp = expected_strings(r)
    pdf = pdf_text(built[eid]["pdf"])
    slides = deck_text(built[eid]["pptx"])
    for key in ("port", "bench", "active"):
        assert exp[key] in pdf, (key, exp[key])
        assert exp[key] in slides[0], (key, exp[key])
    for t in r.tickers:
        assert exp[f"ret_{t}"] in pdf and exp[f"contrib_{t}"] in pdf and exp[f"endw_{t}"] in pdf
    assert r.decision["action"] in pdf and r.decision["action"] in slides[1]
    assert norm(T.pp(r.classes.loc["Equity", "contrib"])) in slides[0] or norm(T.pp(r.classes.loc["Equity", "contrib"])) in pdf
    # slide 2 repeats the weights it shows
    assert norm(T.wt(r.classes.loc["Equity", "end_w"])) in slides[1]


@pytest.mark.parametrize("eid", EDITIONS)
def test_deck_has_three_slides_notes_and_presentation_sized_text(built, results, eid):
    prs = Presentation(built[eid]["pptx"])
    assert len(prs.slides) == 3
    for s in prs.slides:
        assert s.has_notes_slide and len(s.notes_slide.notes_text_frame.text.split()) > 60
        sizes = [run.font.size.pt for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for run in p.runs]
        assert min(sizes) >= 12 and max(sizes) >= 28
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= prs.slide_width + 10 and sh.top + sh.height <= prs.slide_height + 10
        assert sum(1 for sh in s.shapes if sh.shape_type == 13) <= 1
    texts = deck_text(built[eid]["pptx"])
    assert all("Independent educational research sample".upper() in t and "HYPOTHETICAL" in t for t in texts)
    assert "CIBC" not in " ".join(texts)


@pytest.mark.parametrize("eid", EDITIONS)
def test_workbook_structure(built, eid):
    wb = openpyxl.load_workbook(built[eid]["xlsx"])
    assert wb.sheetnames == ["Overview", "Holdings", "Returns", "Contributions", "Assumptions", "Sources"]
    for ws in wb.worksheets:
        assert ws.freeze_panes, ws.title
        assert ws.auto_filter.ref, ws.title
    assert {"Start_Value", "Class_Band", "Hold_Band", "Tol"} <= set(wb.defined_names)
    assert len(wb["Overview"]._charts) == 2
    h = wb["Holdings"]
    assert h["G5"].font.color.rgb.endswith("1D4ED8")  # inputs are blue
    assert str(h["I5"].value).startswith("=") and str(h["K5"].value).startswith("=")  # calculations are formulas
    assert wb["Returns"]["AB6"].value.startswith("=")
    assert "0.00%" in h["I5"].number_format and "pp" in h["L5"].number_format


@pytest.mark.parametrize("eid", EDITIONS)
def test_workbook_values_match_results(built, results, eid):
    r = results[eid]
    ws = openpyxl.load_workbook(built[eid]["xlsx"], data_only=True)
    h, c = ws["Holdings"], ws["Contributions"]
    for i, t in enumerate(r.tickers):
        row = r.holdings.loc[t]
        assert h.cell(5 + i, 1).value == t
        assert h.cell(5 + i, 7).value == pytest.approx(row["begin_w"], abs=1e-12)
        assert h.cell(5 + i, 9).value == pytest.approx(row["total_ret"], abs=1e-12)
        assert h.cell(5 + i, 11).value == pytest.approx(row["end_w"], abs=1e-12)
        assert c.cell(5 + i, 5).value == pytest.approx(row["contrib"] * 100, abs=1e-10)
    assert h["I14"].value == pytest.approx(r.port_return, abs=1e-12)
    assert ws["Returns"]["F12"].value == pytest.approx(r.bench_return, abs=1e-12)
    assert ws["Overview"]["F5"].value == r.decision["action"]
    assert ws["Holdings"]["C15"].value == "OK" and c["B22"].value == "OK"
    assert ws["Holdings"]["H14"].value == pytest.approx(r.start_value)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("eid", EDITIONS)
def test_every_workbook_formula_recalculates_to_its_cached_value(built, tmp_path, eid):
    """Independent recalculation (LibreOffice) must reproduce every cached formula result."""
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path / 'profile'}", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", str(tmp_path), str(built[eid]["xlsx"])], check=True, capture_output=True, timeout=240)
    cached = openpyxl.load_workbook(built[eid]["xlsx"], data_only=True)
    formulas = openpyxl.load_workbook(built[eid]["xlsx"])
    recalc = openpyxl.load_workbook(tmp_path / "workbook.xlsx", data_only=True)
    n = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    a, b = cached[ws.title][cell.coordinate].value, recalc[ws.title][cell.coordinate].value
                    n += 1
                    if isinstance(a, (int, float)):
                        assert b == pytest.approx(a, rel=1e-9, abs=1e-12), (ws.title, cell.coordinate, cell.value)
                    else:
                        assert a == b, (ws.title, cell.coordinate)
    assert n > 400


def test_charts_use_required_visual_encoding(built):
    from PIL import Image

    for eid in EDITIONS:
        for name, path in built[eid]["charts"].items():
            im = Image.open(path)
            assert im.width >= 900 and im.height >= 500
    # portfolio navy solid and benchmark slate dashed are defined once and shared
    import inspect
    from mpr import charts
    src = inspect.getsource(charts.performance_chart)
    assert "T.NAVY" in src and 'ls="-"' in src and "(0, (5, 3))" in src
