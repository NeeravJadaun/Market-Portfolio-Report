"""Three-slide deck with speaker notes. Reuses the report's charts and the shared results object."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import theme as T
from .calc import Results
from .narrative import load_text, source_list


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def _run(p, text, size, color=T.TEXT, bold=False):
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = T.FONT, Pt(size), bold
    r.font.color.rgb = rgb(color)
    return r


def textbox(slide, x, y, w, h, paras, size=18, color=T.TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, space=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.28)))
            pPr.set("indent", str(-Inches(0.28)))
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
        if isinstance(t, tuple):  # (bold lead, rest)
            _run(p, t[0], size, T.NAVY, True)
            _run(p, t[1], size, color, False)
        else:
            _run(p, t, size, color, bold)
    return tb


def card(slide, x, y, w, h, fill=T.SURFACE, line=T.BORDER):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(line)
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    style = s._element.find(qn("p:style"))
    if style is not None:  # theme effectRef would draw a drop shadow
        s._element.remove(style)
    return s


def kpi(slide, x, y, w, h, value, label):
    card(slide, x, y, w, h)
    textbox(slide, x + 0.06, y + 0.12, w - 0.12, 0.5, [value], size=20, color=T.NAVY, bold=True)
    textbox(slide, x + 0.08, y + 0.68, w - 0.16, h - 0.7, [label], size=12, color=T.MUTED)


def frame(slide, r: Results, title: str, idx: int, source_line: str):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(T.CANVAS)
    textbox(slide, 0.5, 0.25, 7.8, 0.35, [f"{T.LABEL.upper()}  |  {T.HYPO.upper()}"], size=12, color=T.NAVY, bold=True)
    textbox(slide, 8.2, 0.25, 4.65, 0.35, [f"Market & Portfolio Report  |  {r.label}"], size=12, color=T.MUTED, align=PP_ALIGN.RIGHT)
    textbox(slide, 0.5, 0.62, 12.35, 0.85, [title], size=30, color=T.NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rule = card(slide, 0.5, 1.5, 12.35, 0.02, fill=T.BORDER, line=T.BORDER)
    textbox(slide, 0.5, 6.95, 11.4, 0.4, [source_line], size=12, color=T.MUTED)
    textbox(slide, 12.0, 6.95, 0.85, 0.4, [f"{idx} / 3"], size=12, color=T.MUTED, align=PP_ALIGN.RIGHT)


def picture(slide, path, x, y, w, h, alt):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def build_deck(r: Results, charts: dict[str, Path], out: Path) -> Path:
    txt = load_text(r)
    sl = txt["slides"]
    refs = {s["id"]: s for s in source_list(r)}
    src = lambda ids: "Sources: " + "; ".join(f"[{refs[i]['n']}] {refs[i]['publisher']}" for i in ids) + ". Full references in the report."
    boc = next(k for k in refs if k.startswith("boc_20"))
    fed = next(k for k in refs if k.startswith("fed_20"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    c = r.classes
    # ---- slide 1: main message
    s = prs.slides.add_slide(blank)
    frame(s, r, sl["s1_title"], 1, src(["prices", "boc_valet"]))
    picture(s, charts["performance_slide"], 0.5, 1.7, 7.6, 4.3, f"Line chart: cumulative total return of the hypothetical portfolio and benchmark, {r.label}")
    textbox(s, 0.5, 6.05, 7.6, 0.8, [f"Cumulative total return since the {r.start:%b %-d} close, CAD. Daily closes."], size=12, color=T.MUTED)
    kx = 8.4
    for i, (v, l) in enumerate([(T.pct(r.port_return), "Portfolio"), (T.pct(r.bench_return), f"Benchmark ({r.benchmark['ticker']})"),
                                (T.pp(r.active_return).replace(T.NBSP + "pp", ""), "Difference (pp)")]):
        kpi(s, kx + i * 1.52, 1.7, 1.42, 1.25, v, l)
    textbox(s, kx, 3.2, 4.45, 3.6, sl["s1_bullets"], size=18, bullets=True, space=12)
    s.notes_slide.notes_text_frame.text = sl["notes"][0]
    # ---- slide 2: portfolio implications
    s = prs.slides.add_slide(blank)
    frame(s, r, sl["s2_title"], 2, src(["prices", "boc_valet"]))
    picture(s, charts["contributions_slide"], 0.5, 1.7, 6.0, 4.3, f"Bar chart: contribution to portfolio return by holding in percentage points, {r.label}")
    textbox(s, 0.5, 6.05, 6.0, 0.8, ["Contribution = start weight x total return, in percentage points."], size=12, color=T.MUTED)
    card(s, 6.9, 1.7, 5.95, 1.3)
    textbox(s, 7.05, 1.78, 2.2, 0.6, [r.decision["action"]], size=28, color=T.NAVY, bold=True)
    textbox(s, 7.05, 2.4, 5.7, 0.6, [f"Equity {T.wt(c.loc['Equity', 'begin_w'])} to {T.wt(c.loc['Equity', 'end_w'])}; fixed income {T.wt(c.loc['Fixed income', 'begin_w'])} to {T.wt(c.loc['Fixed income', 'end_w'])}"], size=14, color=T.MUTED)
    textbox(s, 6.9, 3.25, 5.95, 3.6, sl["s2_bullets"], size=18, bullets=True, space=12)
    s.notes_slide.notes_text_frame.text = sl["notes"][1]
    # ---- slide 3: risks and monitoring
    s = prs.slides.add_slide(blank)
    frame(s, r, sl["s3_title"], 3, src(["prices", "boc_valet", boc, fed]))
    textbox(s, 0.5, 1.7, 6.0, 0.45, ["Downside risks"], size=20, color=T.NAVY, bold=True)
    textbox(s, 0.5, 2.25, 6.0, 2.7, sl["s3_risks"], size=20, bullets=True, space=12)
    textbox(s, 6.9, 1.7, 5.95, 0.45, ["What to monitor"], size=20, color=T.NAVY, bold=True)
    textbox(s, 6.9, 2.25, 5.95, 2.7, sl["s3_monitor"], size=20, bullets=True, space=12)
    card(s, 0.5, 5.35, 12.35, 1.4, fill=T.SURFACE)
    textbox(s, 0.7, 5.4, 11.95, 1.3, [("Counterargument.  ", txt["counter"])], size=17, space=0, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = sl["notes"][2]
    prs.core_properties.title = f"Market & Portfolio Report, {r.label}: deck"
    prs.core_properties.author = T.LABEL
    prs.core_properties.subject = f"{T.HYPO}; not investment advice"
    prs.save(out)
    return out
